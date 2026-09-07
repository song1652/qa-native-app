#!/usr/bin/env python3
"""Export Markdown test cases to an Import Studio-compatible Excel workbook."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = PROJECT_ROOT / "testcases"
DEFAULT_OUTPUT = PROJECT_ROOT / "import" / "qa_native_app_testcases.xlsx"
HEADERS = [
    "No",
    "Scenario ID",
    "대분류",
    "중분류",
    "소분류",
    "테스트 시나리오",
    "전제조건",
    "테스트 절차",
    "기대결과",
    "등급",
    "Locator 힌트",
]


def _section(text: str, heading: str) -> str:
    match = re.search(
        rf"^## {re.escape(heading)}\s*\n(.*?)(?=^## |^---\s*$|\Z)",
        text,
        re.MULTILINE | re.DOTALL,
    )
    return match.group(1).strip() if match else ""


def _subsection(text: str, heading: str) -> str:
    match = re.search(
        rf"^### {re.escape(heading)}\s*\n(.*?)(?=^### |^## |^---\s*$|\Z)",
        text,
        re.MULTILINE | re.DOTALL,
    )
    return match.group(1).strip() if match else ""


def _clean_list(value: str) -> str:
    lines = []
    for line in value.splitlines():
        cleaned = re.sub(r"^\s*(?:[-*]|\d+[.)])\s*", "", line).strip()
        if cleaned:
            lines.append(cleaned)
    return "\n".join(lines)


def parse_testcase(path: Path, platform: str) -> dict[str, str]:
    text = path.read_text(encoding="utf-8")
    title_match = re.search(r"^#\s+TC-(\d+):\s*(.+)$", text, re.MULTILINE)
    if not title_match:
        raise ValueError(f"TC 제목을 찾을 수 없습니다: {path}")

    number, title = title_match.groups()
    priority_match = re.search(r"^priority:\s*(\S+)", text, re.MULTILINE)
    priority = priority_match.group(1).capitalize() if priority_match else "Medium"
    return {
        "scenario_id": f"{platform.upper()}_{int(number):03d}",
        "platform": "Android" if platform == "android" else "iOS",
        "title": title.strip(),
        "purpose": _clean_list(_section(text, "목적")),
        "precondition": _clean_list(_section(text, "전제조건")),
        "steps": _clean_list(_subsection(text, "단계")),
        "expected": _clean_list(_subsection(text, "기대결과")),
        "priority": priority,
        "locators": _clean_list(_subsection(text, "셀렉터 힌트")),
    }


def collect_testcases(source: Path) -> dict[str, list[dict[str, str]]]:
    grouped: dict[str, list[dict[str, str]]] = {"android": [], "ios": []}
    for platform in grouped:
        for path in sorted((source / platform).glob("tc_*.md")):
            grouped[platform].append(parse_testcase(path, platform))
    return grouped


def _write_sheet(workbook: Workbook, name: str, rows: list[dict[str, str]]) -> None:
    sheet = workbook.create_sheet(name)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    title = sheet.cell(1, 1, f"QA 테스트케이스 — {name}")
    title.font = Font(bold=True, color="FFFFFF", size=14)
    title.fill = PatternFill("solid", fgColor="203D4C")
    title.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 28

    for column, header in enumerate(HEADERS, 1):
        cell = sheet.cell(2, column, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2B6170")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for index, case in enumerate(rows, 1):
        values = [
            index,
            case["scenario_id"],
            "네이티브 앱",
            case["platform"],
            case["purpose"],
            case["title"],
            case["precondition"],
            case["steps"],
            case["expected"],
            case["priority"],
            case["locators"],
        ]
        for column, value in enumerate(values, 1):
            cell = sheet.cell(index + 2, column, value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    widths = [7, 18, 14, 12, 38, 34, 42, 48, 44, 10, 48]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:{get_column_letter(len(HEADERS))}{len(rows) + 2}"


def export_testcases(source: Path, output: Path) -> Path:
    grouped = collect_testcases(source)
    workbook = Workbook()
    workbook.remove(workbook.active)
    _write_sheet(workbook, "android", grouped["android"])
    _write_sheet(workbook, "ios", grouped["ios"])
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return output


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    result = export_testcases(args.source.resolve(), args.output.resolve())
    print(result)


if __name__ == "__main__":
    main()
