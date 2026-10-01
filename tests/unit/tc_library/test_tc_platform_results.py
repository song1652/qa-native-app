from __future__ import annotations

from pathlib import Path

import openpyxl

from _tc_model import new_case, platform_key
from _tc_template import analyze_workbook
from _tc_xlsx_import import import_workbook
from _tc_xlsx_export import export_workbook, verify_export
import _tc_library as lib


def _workbook(path: Path, two_rows: bool) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "로그인"
    headings = ["대분류", "기능", "Test Step", "Expected Result"]
    for col, value in enumerate(headings, 1):
        ws.cell(1, col).value = value
    if two_rows:
        ws["E1"] = "환경"
        ws["E2"], ws["F2"] = "And", "iOS"
        start = 3
    else:
        ws["E1"], ws["F1"] = "Android", "iOS"
        start = 2
    for offset, (a, i) in enumerate((("Pass", "Fail"), ("NA", "Pass"))):
        row = start + offset
        ws.cell(row, 1).value = "로그인"
        ws.cell(row, 2).value = f"경우 {offset}"
        ws.cell(row, 3).value = "1. 버튼 누르기"
        ws.cell(row, 4).value = "화면 표시"
        ws.cell(row, 5).value = a
        ws.cell(row, 6).value = i
    wb.save(path)
    wb.close()
    return path


def test_platform_alias_and_default_case():
    assert [platform_key(s) for s in (" And ", "Android", "AOS", "안드로이드", "IOS", "아이폰")] == [
        "android", "android", "android", "android", "ios", "ios"]
    case = new_case()
    assert case["platforms"] == ["android", "ios"] and case["results"] == {}


def test_two_row_result_columns_roundtrip(tmp_path):
    source = _workbook(tmp_path / "two.xlsx", True)
    profiles = analyze_workbook(source)
    cases = import_workbook(source, profiles, ["로그인"], {"로그인": "LOG"})
    assert cases[0]["results"] == {"android": "pass", "ios": "fail"}
    assert cases[0]["execution_result"] == "fail"
    assert cases[1]["platforms"] == ["ios"]
    assert cases[1]["results"] == {"android": "na", "ios": "pass"}
    out = tmp_path / "out.xlsx"
    export_workbook(source, profiles, {"로그인": cases}, out)
    again = import_workbook(out, analyze_workbook(out), ["로그인"], {"로그인": "LOG"})
    assert [(c["results"], c["platforms"]) for c in again] == [
        (c["results"], c["platforms"]) for c in cases]


def test_one_row_result_columns_roundtrip(tmp_path):
    source = _workbook(tmp_path / "one.xlsx", False)
    profiles = analyze_workbook(source)
    cases = import_workbook(source, profiles, ["로그인"], {"로그인": "LOG"})
    assert profiles["로그인"].result_columns == {"Android": 5, "iOS": 6}
    assert cases[0]["results"] == {"android": "pass", "ios": "fail"}
    out = tmp_path / "out.xlsx"
    export_workbook(source, profiles, {"로그인": cases}, out)
    again = import_workbook(out, analyze_workbook(out), ["로그인"], {"로그인": "LOG"})
    assert again[0]["results"] == cases[0]["results"]


def test_export_verifier_checks_each_platform_result(tmp_path):
    source = _workbook(tmp_path / "one.xlsx", False)
    profiles = analyze_workbook(source)
    cases = import_workbook(source, profiles, ["로그인"], {"로그인": "LOG"})
    out = tmp_path / "out.xlsx"
    export_workbook(source, profiles, {"로그인": cases}, out)
    book = openpyxl.load_workbook(out)
    book.active["E2"], book.active["F2"] = "Fail", "Pass"
    book.save(out)
    book.close()
    assert any(c["level"] == "error" and c["code"] == "ROUNDTRIP"
               for c in verify_export(out, profiles, {"로그인": cases}))


def test_ios_only_draft_exports_android_as_na(tmp_path):
    source = _workbook(tmp_path / "one.xlsx", False)
    profiles = analyze_workbook(source)
    case = new_case(case_id="APP_0001", sheet="로그인", path=["로그인", "", ""], feature="iOS 전용",
                    steps=["실행"], expected="표시", platforms=["ios"])
    out = tmp_path / "out.xlsx"
    export_workbook(source, profiles, {"로그인": [case]}, out)
    book = openpyxl.load_workbook(out)
    assert (book.active["E2"].value, book.active["F2"].value) == ("NA", None)
    book.close()
    assert not [c for c in verify_export(out, profiles, {"로그인": [case]}) if c["level"] == "error"]


def test_patch_recomputes_merged_result(library_dir):
    lib.import_cases("앱", ["로그인"], [new_case(case_id="APP_0001", sheet="로그인", feature="로그인",
        path=["로그인", "", ""], steps=["누른다"], expected="표시", results={"android": "pass", "ios": "pass"})], "t")
    saved = lib.patch_case("앱", "APP_0001", 1, {"results": {"android": "pass", "ios": "fail"}}, "t")
    assert saved["execution_result"] == "fail"
    assert [c["case_id"] for c in lib.filter_cases(lib.load_cases("앱"), {"mismatch": "1"})] == ["APP_0001"]
    assert lib.list_suites()[0]["kind"] == "app"


def test_mismatch_includes_one_unrun_platform_but_excludes_exclusive_case():
    shared = new_case(case_id="APP_0001", results={"android": "pass", "ios": ""})
    exclusive = new_case(case_id="APP_0002", platforms=["ios"], results={"android": "na", "ios": "pass"})
    assert [c["case_id"] for c in lib.filter_cases([shared, exclusive], {"mismatch": "1"})] == ["APP_0001"]


def test_bulk_result_changes_only_chosen_platform(library_dir):
    lib.import_cases("앱", ["로그인"], [new_case(case_id="APP_0001", sheet="로그인", feature="로그인",
        path=["로그인", "", ""], steps=["누른다"], expected="표시", results={"android": "pass", "ios": "pass"})], "t")
    result = lib.bulk_set_result("앱", [{"case_id": "APP_0001", "rev": 1}], "ios", "fail", "t")
    assert result["updated"] == ["APP_0001"]
    saved = lib.get_case("앱", "APP_0001")
    assert saved["results"] == {"android": "pass", "ios": "fail"}
    assert saved["execution_result"] == "fail"
