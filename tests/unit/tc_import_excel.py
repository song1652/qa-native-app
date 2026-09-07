import importlib.util
from pathlib import Path

from openpyxl import Workbook


MODULE_PATH = Path(__file__).parents[2] / "scripts" / "import_excel.py"
SPEC = importlib.util.spec_from_file_location("import_excel", MODULE_PATH)
assert SPEC and SPEC.loader
import_excel = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(import_excel)


def test_convert_sheet_applies_selected_column_mappings(tmp_path):
    source = tmp_path / "cases.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "login"
    sheet.append(["unused", "Test Case", "ID", "Title", "Precondition", "Steps", "Expected"])
    sheet.append(["", "ignored", "LOGIN_007", "로그인 성공", "로그인 화면", "아이디 입력", "홈 노출"])
    workbook.save(source)

    created = import_excel.convert_sheet(
        source,
        "login",
        tmp_path / "testcases",
        ["android"],
        {
            "tc_id": "C열",
            "title": "D열",
            "precondition": "E열",
            "steps": "F열",
            "expected": "G열",
        },
    )

    assert len(created) == 1
    assert created[0].parent.name == "login"
    assert created[0].parent.parent.name == "android"
    assert created[0].name.startswith("tc_007_")
    content = created[0].read_text(encoding="utf-8")
    assert "TC-007: 로그인 성공" in content
    assert "- 로그인 화면" in content
    assert "1. 아이디 입력" in content
    assert "- 홈 노출" in content


def test_convert_sheet_splits_android_and_ios_outputs(tmp_path):
    source = tmp_path / "cases.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "login"
    sheet.append(["", "Test Case", "", "", "", "Title", "Pre", "Steps", "Expected"])
    sheet.append(["", "LOGIN_001", "", "", "", "로그인", "준비", "실행", "완료"])
    workbook.save(source)

    created = import_excel.convert_sheet(
        source, "login", tmp_path / "testcases", ["android", "ios"]
    )

    assert {path.parent.parent.name for path in created} == {"android", "ios"}
    android = next(path for path in created if path.parent.parent.name == "android")
    ios = next(path for path in created if path.parent.parent.name == "ios")
    assert "## 플랫폼\n- Android" in android.read_text(encoding="utf-8")
    assert "## 플랫폼\n- iOS" in ios.read_text(encoding="utf-8")


def test_convert_sheet_skip_conflict_preserves_existing_markdown(tmp_path):
    source = tmp_path / "cases.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "login"
    sheet.append(["", "Test Case", "", "", "", "Title", "Pre", "Steps", "Expected"])
    sheet.append(["", "LOGIN_001", "", "", "", "로그인", "준비", "실행", "완료"])
    workbook.save(source)

    created = import_excel.convert_sheet(
        source, "login", tmp_path / "testcases", ["android"], policy="overwrite"
    )
    created[0].write_text("사용자가 수정한 기존 내용", encoding="utf-8")

    skipped = import_excel.convert_sheet(
        source, "login", tmp_path / "testcases", ["android"], policy="skip-conflict"
    )

    assert skipped == []
    assert created[0].read_text(encoding="utf-8") == "사용자가 수정한 기존 내용"


def test_platform_named_sheet_only_writes_to_matching_platform_root(tmp_path):
    source = tmp_path / "cases.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "android"
    sheet.append(["", "Test Case", "", "", "", "Title", "Pre", "Steps", "Expected"])
    sheet.append(["", "ANDROID_001", "", "", "", "설정 화면", "준비", "실행", "완료"])
    workbook.save(source)

    created = import_excel.convert_sheet(
        source, "android", tmp_path / "testcases", ["android", "ios"]
    )

    assert len(created) == 1
    assert created[0].parent == tmp_path / "testcases" / "android"
    assert not (tmp_path / "testcases" / "android" / "android").exists()
    assert not (tmp_path / "testcases" / "ios").exists()


def test_platform_named_sheet_is_skipped_when_platform_is_not_selected(tmp_path):
    source = tmp_path / "cases.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ios"
    sheet.append(["", "Test Case", "", "", "", "Title", "Pre", "Steps", "Expected"])
    sheet.append(["", "IOS_001", "", "", "", "설정 화면", "준비", "실행", "완료"])
    workbook.save(source)

    created = import_excel.convert_sheet(
        source, "ios", tmp_path / "testcases", ["android"]
    )

    assert created == []
    assert not (tmp_path / "testcases").exists()


def test_preview_sheets_reports_added_and_updates_when_mapping_changes(tmp_path):
    source = tmp_path / "cases.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "android"
    sheet.append(["", "Test Case", "", "", "", "Title", "Pre", "Steps", "Expected"])
    sheet.append(["", "ANDROID_001", "", "", "", "설정 화면", "준비", "실행", "완료"])
    workbook.save(source)

    preview = import_excel.preview_sheets(
        source, ["android"], tmp_path / "testcases", ["android", "ios"]
    )

    assert preview["summary"] == {
        "added": 1,
        "updated": 0,
        "conflict": 0,
        "error": 0,
        "same": 0,
    }
    assert preview["rows"][0]["tc_id"] == "ANDROID_001"
    assert preview["rows"][0]["platform"] == "android"

    import_excel.convert_sheet(
        source, "android", tmp_path / "testcases", ["android"], policy="overwrite"
    )
    unchanged = import_excel.preview_sheets(
        source, ["android"], tmp_path / "testcases", ["android"]
    )
    assert unchanged["summary"]["same"] == 1
    assert unchanged["summary"]["updated"] == 0

    invalid = import_excel.preview_sheets(
        source,
        ["android"],
        tmp_path / "testcases",
        ["android"],
        {"title": "Z열"},
    )

    assert invalid["summary"]["error"] == 1
    assert invalid["rows"][0]["status"] == "error"


def test_preview_sheets_counts_all_actionable_statuses(tmp_path):
    source = tmp_path / "cases.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "android"
    sheet.append(["", "Test Case", "", "", "", "Title", "Pre", "Steps", "Expected"])
    sheet.append(["", "ANDROID_001", "", "", "", "새 파일", "준비", "실행", "완료"])
    sheet.append(["", "ANDROID_002", "", "", "", "변경 파일", "준비", "실행", "완료"])
    sheet.append(["", "ANDROID_003", "", "", "", "중복 파일", "준비", "실행", "완료"])
    sheet.append(["", "ANDROID_003", "", "", "", "중복 파일", "준비", "실행", "완료"])
    sheet.append(["", "ANDROID_004", "", "", "", "", "준비", "실행", "완료"])
    workbook.save(source)

    seed = tmp_path / "seed.xlsx"
    seed_workbook = Workbook()
    seed_sheet = seed_workbook.active
    seed_sheet.title = "android"
    seed_sheet.append(["", "Test Case", "", "", "", "Title", "Pre", "Steps", "Expected"])
    seed_sheet.append(["", "ANDROID_002", "", "", "", "변경 파일", "준비", "실행", "완료"])
    seed_workbook.save(seed)
    existing = import_excel.convert_sheet(seed, "android", tmp_path / "testcases", ["android"])
    existing[0].write_text("기존의 다른 내용", encoding="utf-8")

    preview = import_excel.preview_sheets(
        source, ["android"], tmp_path / "testcases", ["android"]
    )

    assert preview["summary"] == {
        "added": 2,
        "updated": 1,
        "conflict": 1,
        "error": 1,
        "same": 0,
    }
